#!/usr/bin/env python3
"""eng-report engine: engineering house style for Word documents, in your own branding.

  brand                                 save a brand profile (company, logo, colours)
  restyle IN.docx OUT.docx [options]    re-skin an existing document in place of its styles
  build   IN.md   OUT.docx [options]    build a document from Markdown (+ front matter)
  check   FILE.docx                     list house-rule problems (no changes)

Brand profile: brand --company "Acme Controls" [--tagline ...] [--logo logo.png]
  [--accent 1F4E79] [--head 1A1A1A] [--link 1F4E79] [--scope project|user]
  restyle/build read --brand-file, else ./.eng-report/brand.json, else ~/.eng-report/brand.json;
  with none they stop and ask for one. --plain runs without a profile (no logo, neutral accent).
Options (restyle/build): --title  --subtitle  --doc-type  --doc-no  --rev  --date "October 2026"
  --author  --classification "Internal"  --cover/--no-cover  --numbered/--no-numbered
  --page letter|a4|keep
Markdown extras: front matter (--- key: value ---), notices as "> **WARNING:** hazard. text",
"> **NOTE:** text"; tables in GitHub pipe syntax; "[[revision-history]]" and "[[sign-off]]"
placeholders; captions as lines starting "Table: " (above a table) or "Figure: ".
"""
import argparse, json, os, re, shutil, sys
from datetime import date
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Emu

INK = '1A1A1A'; CHARCOAL = '3E3E3E'; GREY = '6B6B6B'; RULE = 'BFBFBF'; PALE = 'F1F1F1'
# accent = rules, bands, cover block (never text: a light accent on white fails contrast)
PLAIN = dict(company='', tagline='', logo='', accent='1F4E79', head=INK, link='1F4E79')
PROFILE_DIR = '.eng-report'; PROFILE = 'brand.json'
LOGO_TYPES = ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff')
LINK = PLAIN['link']

def hexcolor(v, name):
    v = (v or '').strip().lstrip('#').upper()
    if not re.fullmatch(r'[0-9A-F]{6}', v): sys.exit(f'--{name}: expected a 6-digit hex colour like 1F4E79, got "{v}"')
    return v

def find_profile(explicit=None):
    for p in ([explicit] if explicit else [os.path.join(os.getcwd(), PROFILE_DIR, PROFILE),
                                           os.path.join(os.path.expanduser('~'), PROFILE_DIR, PROFILE)]):
        if p and os.path.isfile(p): return p
    return None

def load_brand(a):
    """Brand dict for this run. The logo path is resolved next to the profile file."""
    global LINK
    if a.plain: b = dict(PLAIN)
    else:
        if a.brand_file and not os.path.isfile(a.brand_file): sys.exit(f'--brand-file: no such file: {a.brand_file}')
        path = find_profile(a.brand_file)
        if not path:
            sys.exit('NO BRAND PROFILE. Ask the user for their company name, logo file (PNG or JPG) and brand '
                     'colours, then run:  eng_report.py brand --company "..." --logo <file> --accent <hex>\n'
                     'Or pass --plain for an unbranded document (no logo, neutral accent).')
        with open(path, encoding='utf-8') as f: b = dict(PLAIN, **json.load(f))
        if not str(b.get('company') or '').strip(): sys.exit(f'brand profile {path}: "company" is missing')
        if b.get('logo'):
            b['logo'] = os.path.join(os.path.dirname(os.path.abspath(path)), b['logo'])
            if not os.path.isfile(b['logo']): sys.exit(f"brand profile {path}: logo file not found: {b['logo']}")
        print('brand profile:', path)
    for k in ('accent', 'head', 'link'): b[k] = hexcolor(b[k], k)
    LINK = b['link']
    return b

def save_brand(a):
    if not a.company: sys.exit('brand: --company is required')
    root = os.path.expanduser('~') if a.scope == 'user' else os.getcwd()
    d = os.path.join(root, PROFILE_DIR); os.makedirs(d, exist_ok=True)
    b = dict(company=a.company, tagline=a.tagline or '', logo='',
             accent=hexcolor(a.accent or PLAIN['accent'], 'accent'), head=hexcolor(a.head or PLAIN['head'], 'head'),
             link=hexcolor(a.link or a.accent or PLAIN['link'], 'link'))
    path = os.path.join(d, PROFILE)
    if not a.logo and not a.no_logo and os.path.isfile(path):
        try:
            with open(path, encoding='utf-8') as f: old = json.load(f).get('logo') or ''
            if old and os.path.isfile(os.path.join(d, old)): b['logo'] = old; print('kept the stored logo:', old)
        except Exception: pass
    if a.logo:
        ext = os.path.splitext(a.logo)[1].lower()
        if ext not in LOGO_TYPES: sys.exit(f'brand: logo must be one of {", ".join(LOGO_TYPES)} (Word cannot embed {ext or "that file"}; export a PNG)')
        if not os.path.isfile(a.logo): sys.exit(f'brand: logo file not found: {a.logo}')
        b['logo'] = 'logo' + ext; dest = os.path.join(d, b['logo'])
        if os.path.abspath(a.logo) != os.path.abspath(dest): shutil.copyfile(a.logo, dest)
    with open(path, 'w', encoding='utf-8', newline='\n') as f: json.dump(b, f, indent=1); f.write('\n')
    print('saved brand profile:', path)
    if not b['logo']: print('no logo given: the header and cover will show the company name as text')

def brand_mark(paragraph, b, height, size):
    """Logo when the profile has one, otherwise the company name set in type."""
    if b['logo']: paragraph.add_run().add_picture(b['logo'], height=Inches(height))
    elif b['company']: set_font(paragraph.add_run(b['company']), size, True, b['head'])

# ANSI Z535.6 signal-word panels: (fill, text colour, italic, alert symbol)
NOTICES = {
    'DANGER':  ('C8102E', 'FFFFFF', False, True),
    'WARNING': ('FF7900', '000000', False, True),
    'CAUTION': ('FFD100', '000000', False, True),
    'NOTICE':  ('005EB8', 'FFFFFF', True,  False),
}
INFO = {'NOTE', 'TIP', 'IMPORTANT'}
FONT = 'Arial'

# ---------- low-level XML helpers ----------
def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn('w:shd')): tcPr.remove(e)
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexfill)
    tcPr.append(s)

def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for k, v in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    for e in tcPr.findall(qn('w:tcMar')): tcPr.remove(e)
    tcPr.append(m)

def table_borders(table, color=RULE, size=4, inside=True, outer=True):
    tblPr = table._tbl.tblPr
    for e in tblPr.findall(qn('w:tblBorders')): tblPr.remove(e)
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        on = (inside if side.startswith('inside') else outer)
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single' if on else 'nil'); e.set(qn('w:sz'), str(size)); e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)

def para_border(p, side='bottom', color='000000', size=12, space=4):
    pPr = p._p.get_or_add_pPr(); bdr = pPr.find(qn('w:pBdr'))
    if bdr is None: bdr = OxmlElement('w:pBdr'); pPr.append(bdr)
    for e in bdr.findall(qn(f'w:{side}')): bdr.remove(e)
    e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(size))
    e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color); bdr.append(e)

def style_border(style, color, size=8, space=2):
    pPr = style.element.get_or_add_pPr()
    for e in pPr.findall(qn('w:pBdr')): pPr.remove(e)
    bdr = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(size)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
    bdr.append(e); pPr.append(bdr)

def field(run, instr):
    for kind, text in (('begin', None), (None, instr), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); run._r.append(f)
        elif text == instr:
            t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = f' {instr} '; run._r.append(t)
        else:
            t = OxmlElement('w:t'); t.text = text; run._r.append(t)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)

def cant_split(row):
    trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))

def set_font(run, size=None, bold=None, color=None, italic=None, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr(); rf = rPr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color)

# ---------- theme ----------
def apply_styles(doc, b, numbered=True):
    st = doc.styles
    def base(name, size, bold=False, color=INK, before=0, after=6, keep=False, italic=False):
        try: s = st[name]
        except KeyError: s = st.add_style(name, 1)
        s.font.name = FONT; s.font.size = Pt(size); s.font.bold = bold; s.font.italic = italic
        s.font.color.rgb = RGBColor.from_string(color)
        rPr = s.element.get_or_add_rPr(); rf = rPr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), FONT)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        pf = s.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
        pf.keep_with_next = keep
        return s
    n = base('Normal', 10.5, after=6); n.paragraph_format.line_spacing = 1.15
    h1 = base('Heading 1', 15, True, b['head'], before=18, after=8, keep=True)
    style_border(h1, b['accent'], size=12, space=3)
    h1.paragraph_format.page_break_before = False
    base('Heading 2', 12.5, True, b['head'], before=14, after=5, keep=True)
    base('Heading 3', 11, True, CHARCOAL, before=10, after=4, keep=True)
    base('Heading 4', 10.5, True, CHARCOAL, before=8, after=3, keep=True)
    t = base('Title', 26, True, INK, after=4)
    for e in t.element.get_or_add_pPr().findall(qn('w:pBdr')): t.element.get_or_add_pPr().remove(e)
    _nb = OxmlElement('w:pBdr'); _b = OxmlElement('w:bottom'); _b.set(qn('w:val'), 'nil'); _nb.append(_b); t.element.get_or_add_pPr().append(_nb)
    base('Subtitle', 13, False, CHARCOAL, after=12)
    base('Caption', 9, True, CHARCOAL, before=6, after=4, keep=True)
    for lst in ('List Bullet', 'List Number', 'List Bullet 2', 'List Number 2'):
        try: base(lst, 10.5, after=3)
        except Exception: pass
    if numbered: heading_numbering(doc)
    return b

def heading_numbering(doc):
    """Decimal 1 / 1.1 / 1.1.1 numbering bound to Heading 1-3 (Siemens / E+H scheme)."""
    numbering = doc.part.numbering_part.element
    aid = 9000 + len(numbering.findall(qn('w:abstractNum')))
    a = OxmlElement('w:abstractNum'); a.set(qn('w:abstractNumId'), str(aid))
    ml = OxmlElement('w:multiLevelType'); ml.set(qn('w:val'), 'multilevel'); a.append(ml)
    for i in range(3):
        lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), str(i))
        for tag, val in (('w:start', '1'), ('w:numFmt', 'decimal'),
                         ('w:pStyle', f'Heading{i+1}'),
                         ('w:lvlText', '.'.join(f'%{k+1}' for k in range(i+1))), ('w:lvlJc', 'left')):
            e = OxmlElement(tag); e.set(qn('w:val'), val); lvl.append(e)
        pPr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), str(567 + 142 * i)); ind.set(qn('w:hanging'), str(567 + 142 * i)); pPr.append(ind); lvl.append(pPr)
        a.append(lvl)
    first_num = numbering.find(qn('w:num'))
    if first_num is not None: first_num.addprevious(a)
    else: numbering.append(a)
    nid = 9000 + len(numbering.findall(qn('w:num')))
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(nid))
    ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(aid)); num.append(ref); numbering.append(num)
    for i in range(3):
        s = doc.styles[f'Heading {i+1}']; pPr = s.element.get_or_add_pPr()
        for e in pPr.findall(qn('w:numPr')): pPr.remove(e)
        numPr = OxmlElement('w:numPr'); il = OxmlElement('w:ilvl'); il.set(qn('w:val'), str(i))
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(nid)); numPr.append(il); numPr.append(ni)
        pPr.insert(0, numPr)

def page_setup(doc, page):
    for s in doc.sections:
        if page == 'a4': s.page_width, s.page_height = Emu(7560310), Emu(10692130)
        elif page == 'letter': s.page_width, s.page_height = Inches(8.5), Inches(11)
        if s.orientation == WD_ORIENT.LANDSCAPE and s.page_width < s.page_height:
            s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = Inches(0.95); s.bottom_margin = Inches(0.8)
        s.header_distance = Inches(0.4); s.footer_distance = Inches(0.4)

def content_width(sec):
    return sec.page_width - sec.left_margin - sec.right_margin

def header_footer(doc, meta, b, cover):
    for i, s in enumerate(doc.sections):
        s.different_first_page_header_footer = bool(cover and i == 0)
        hdr = s.header; hdr.is_linked_to_previous = False
        for p in list(hdr.paragraphs)[1:]: p._p.getparent().remove(p._p)
        for t in list(hdr.tables): t._tbl.getparent().remove(t._tbl)
        p = hdr.paragraphs[0]; p.text = ''
        w = content_width(s)
        tbl = hdr.add_table(rows=1, cols=2, width=w); table_borders(tbl, inside=False, outer=False)
        lc, rc = tbl.rows[0].cells
        lc.width = int(w * 0.4); rc.width = int(w * 0.6)
        brand_mark(lc.paragraphs[0], b, 0.32, 11)
        rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(rp.add_run(meta.get('title', '')), 8.5, True, CHARCOAL)
        if meta.get('doc_type'):
            rp2 = rc.add_paragraph(); rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_font(rp2.add_run(meta['doc_type']), 8, False, GREY)
        p._p.getparent().remove(p._p)
        rule = hdr.add_paragraph(); para_border(rule, 'bottom', b['accent'], 12, 1)
        rule.paragraph_format.space_after = Pt(0); set_font(rule.add_run(''), 2)

        ftr = s.footer; ftr.is_linked_to_previous = False
        for t in list(ftr.tables): t._tbl.getparent().remove(t._tbl)
        for p in list(ftr.paragraphs)[1:]: p._p.getparent().remove(p._p)
        fp0 = ftr.paragraphs[0]; fp0.text = ''; para_border(fp0, 'bottom', RULE, 4, 2)
        fp0.paragraph_format.space_after = Pt(2); set_font(fp0.add_run(''), 2)
        bits = [x for x in (meta.get('doc_no'), f"Rev {meta['rev']}" if meta.get('rev') else None,
                            meta.get('date'), meta.get('classification')) if x]
        left = '  |  '.join(bits) if bits else b['company']
        w = content_width(s); ft = ftr.add_table(rows=1, cols=2, width=w); table_borders(ft, inside=False, outer=False)
        lc, rc = ft.rows[0].cells; lc.width = int(w * 0.7); rc.width = int(w * 0.3)
        set_font(lc.paragraphs[0].add_run(left), 8, False, GREY)
        fp = rc.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run('Page '); set_font(r, 8, False, GREY)
        r = fp.add_run(); set_font(r, 8, True, CHARCOAL); field(r, 'PAGE')
        r = fp.add_run(' of '); set_font(r, 8, False, GREY)
        r = fp.add_run(); set_font(r, 8, True, CHARCOAL); field(r, 'NUMPAGES')
        if b['company']:
            fp2 = ftr.add_paragraph(); set_font(fp2.add_run(f"© {date.today().year} {b['company']}"
                                             + (f" · {b['tagline']}" if b['tagline'] else '')), 7, False, GREY)

def cover_page(doc, meta, b):
    """Insert a cover at the very start: logo, accent band, type, title, id table."""
    body = doc.element.body; first = body[0]
    def add(el): first.addprevious(el); return el
    p = doc.add_paragraph(); brand_mark(p, b, 0.85, 22)
    p.paragraph_format.space_after = Pt(60); add(p._p)
    band = doc.add_paragraph(); para_border(band, 'top', b['accent'], 48, 0); band.paragraph_format.space_after = Pt(18); add(band._p)
    if meta.get('doc_type'):
        p = doc.add_paragraph(); set_font(p.add_run(meta['doc_type'].upper()), 10, True, GREY)
        p.paragraph_format.space_after = Pt(6); add(p._p)
    p = doc.add_paragraph(style='Title'); p.add_run(meta.get('title', 'Untitled')); add(p._p)
    if meta.get('subtitle'):
        p = doc.add_paragraph(style='Subtitle'); p.add_run(meta['subtitle']); add(p._p)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(150); add(sp._p)
    rows = [(k, v) for k, v in (('Document number', meta.get('doc_no')), ('Revision', meta.get('rev')),
            ('Date', meta.get('date')), ('Prepared by', meta.get('author')),
            ('Classification', meta.get('classification'))) if v]
    if rows:
        t = doc.add_table(rows=len(rows), cols=2); table_borders(t, RULE, 4, inside=True, outer=False)
        for i, (k, v) in enumerate(rows):
            a, c = t.rows[i].cells; a.text = ''; c.text = ''
            set_font(a.paragraphs[0].add_run(k), 9, True, CHARCOAL); set_font(c.paragraphs[0].add_run(str(v)), 9, False, INK)
            a.width = Inches(1.8); c.width = Inches(3.2)
        add(t._tbl)
    if b['company']:
        co = doc.add_paragraph(); set_font(co.add_run(b['company'] + (f" · {b['tagline']}" if b['tagline'] else '')), 9, True, CHARCOAL)
        co.paragraph_format.space_before = Pt(18); add(co._p)
    br = doc.add_paragraph(); br.add_run().add_break(WD_BREAK.PAGE); add(br._p)

def notice_block(doc, kind, text, anchor=None):
    """ANSI-style panel: coloured signal-word cell + message cell. Info kinds are a grey bar."""
    kind = kind.upper()
    t = doc.add_table(rows=1, cols=2) if kind in NOTICES else doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; cant_split(t.rows[0])
    sec = doc.sections[-1]; w = content_width(sec)
    if kind in NOTICES:
        fill, fg, ital, alert = NOTICES[kind]
        table_borders(t, fill, 8, inside=False, outer=True)
        sc, mc = t.rows[0].cells; sc.width = Inches(1.25); mc.width = w - Inches(1.25)
        shade(sc, fill); cell_margins(sc); cell_margins(mc, left=140)
        sp = sc.paragraphs[0]; sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(sp.add_run(('⚠ ' if alert else '') + kind), 10, True, fg, italic=ital)
        target = mc
    else:
        table_borders(t, PALE, 4, inside=False, outer=False)
        c = t.rows[0].cells[0]; shade(c, PALE); cell_margins(c, left=160)
        tcPr = c._tc.get_or_add_tcPr(); bd = OxmlElement('w:tcBorders'); l = OxmlElement('w:start')
        l.set(qn('w:val'), 'single'); l.set(qn('w:sz'), '24'); l.set(qn('w:color'), CHARCOAL); bd.append(l); tcPr.append(bd)
        target = c
        p0 = c.paragraphs[0]; set_font(p0.add_run(kind.title() if kind != 'IMPORTANT' else 'IMPORTANT'), 9, True, CHARCOAL)
    m = re.match(r'\s*\*{0,2}([^.*]{3,90}?)\*{0,2}\.\s+(.*)', text, re.S)
    p = target.paragraphs[0] if kind in NOTICES else target.add_paragraph()
    if kind in NOTICES and m:
        set_font(p.add_run(m.group(1).strip()), 10, True, INK)
        p2 = target.add_paragraph(); add_inline(p2, m.group(2).strip())
    else:
        add_inline(p, text.strip())
    for cp in target.paragraphs: cp.paragraph_format.space_after = Pt(2)
    gap = doc.add_paragraph(); gap.paragraph_format.space_after = Pt(2)
    if anchor is not None:
        anchor.addprevious(t._tbl); anchor.addprevious(gap._p)
    return t

COVER_LABELS = {'Document number', 'Revision', 'Date', 'Prepared by', 'Classification'}
def is_engine_table(t):
    """A notice panel or the cover id table from an earlier run: restyle must leave it alone."""
    try:
        first = t.rows[0].cells[0]
        shd = first._tc.get_or_add_tcPr().find(qn('w:shd'))
        fill = (shd.get(qn('w:fill')) or '').upper() if shd is not None else ''
        if fill in {v[0] for v in NOTICES.values()} | {PALE}: return True
        return len(t.columns) == 2 and all(r.cells[0].text.strip() in COVER_LABELS for r in t.rows)
    except Exception:
        return False

def style_table(t, b, header=True):
    table_borders(t, RULE, 4, inside=True, outer=True)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(t.rows):
        for c in row.cells:
            cell_margins(c, 50, 50, 90, 90)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
                for r in p.runs:
                    set_font(r, 9.5, True if (header and ri == 0) else None,
                             'FFFFFF' if (header and ri == 0) else None)
            if header and ri == 0: shade(c, CHARCOAL if b['head'] == INK else b['head'])
    if header and len(t.rows) > 1: repeat_header(t.rows[0])

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))')
def add_inline(p, text):
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'): r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; set_font(r, 9.5, name='Consolas'); continue
        elif part.startswith('['):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part); r = p.add_run(f'{m.group(1)}'); r.font.color.rgb = RGBColor.from_string(LINK); r.underline = True
        elif part.startswith('*'): r = p.add_run(part[1:-1]); r.italic = True
        else: r = p.add_run(part)
        set_font(r)

# ---------- restyle an existing document ----------
TYPED_NO = r'^\s*(?:\d+(?:\.\d+)+\.?\s+|\d+\.\s+|\d{1,2}\s+(?=[A-Z][a-z]))'
SIGNAL_RE = re.compile(r'^\s*(?:⚠\s*)?\**\s*(DANGER|WARNING|CAUTION|NOTICE|NOTE|Note|TIP|Tip|IMPORTANT|Important)\s*\**\s*(?::|(?<=\s)[\-–—](?=\s))\s*(.+)$', re.S)
def strip_direct_fonts(doc):
    def runs():
        for p in doc.paragraphs: yield from p.runs
        for t in doc.tables:
            for row in t.rows:
                for c in row.cells:
                    for p in c.paragraphs: yield from p.runs
    for r in runs():
        rPr = r._r.rPr
        if rPr is None: continue
        for tag in ('w:rFonts', 'w:sz', 'w:szCs', 'w:color'):
            for e in rPr.findall(qn(tag)): rPr.remove(e)

def restyle(inp, out, meta, a):
    doc = Document(inp)
    meta.setdefault('date', date.today().strftime('%B %Y'))
    page_setup(doc, a.page)
    b = apply_styles(doc, load_brand(a), a.numbered)
    strip_direct_fonts(doc)
    # numbered headings: drop typed numbers ("1.2 Scope" -> "Scope") so auto-numbering doesn't double up
    if a.numbered:
        for p in doc.paragraphs:
            if p.style.name in ('Heading 1', 'Heading 2', 'Heading 3') and p.runs:
                m = re.match(TYPED_NO + r'|^\s*[A-Z]\.\s+', p.text)
                if m:
                    cut = len(m.group(0))
                    for r in p.runs:
                        if cut <= 0: break
                        n = min(cut, len(r.text)); r.text = r.text[n:]; cut -= n
    for t in doc.tables:
        if not is_engine_table(t): style_table(t, b)
    for p in list(doc.paragraphs):
        m = SIGNAL_RE.match(p.text)
        if m and p.style.name not in ('Heading 1', 'Heading 2', 'Heading 3', 'Title'):
            notice_block(doc, m.group(1), m.group(2), anchor=p._p); p._p.getparent().remove(p._p)
    if a.cover: cover_page(doc, meta, b)
    header_footer(doc, meta, b, a.cover)
    finish(doc, meta, b, out)

# ---------- build from markdown ----------
def parse_front(md):
    meta = {}
    if md.startswith('---'):
        end = md.find('\n---', 3)
        if end > 0:
            for line in md[3:end].splitlines():
                if ':' in line:
                    k, v = line.split(':', 1); meta[k.strip().lower().replace('-', '_').replace(' ', '_')] = v.strip().strip('"')
            md = md[end + 4:]
    return meta, md

def revision_table(doc, rows):
    t = doc.add_table(rows=1, cols=4)
    for c, h in zip(t.rows[0].cells, ('Rev', 'Date', 'Description of change', 'By')): c.text = h
    for r in rows or [('A', date.today().strftime('%B %Y'), 'Initial release', '')]:
        cells = t.add_row().cells
        for c, v in zip(cells, r): c.text = v
    return t

def signoff_table(doc):
    t = doc.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ('Role', 'Name', 'Company', 'Signature', 'Date')): c.text = h
    for role in ('Prepared by', 'Reviewed by', 'Approved by', 'Customer witness'):
        cells = t.add_row().cells; cells[0].text = role
        t.rows[-1].height = Inches(0.45)
    return t

def build(inp, out, meta, a):
    md = open(inp, encoding='utf-8').read().replace('\r\n', '\n')
    fm, md = parse_front(md)
    for k, v in fm.items():
        if v: meta.setdefault(k, v)
    meta.setdefault('date', date.today().strftime('%B %Y'))
    if not meta.get('title'):
        m = re.search(r'^#\s+(.+)$', md, re.M)
        if m: meta['title'] = m.group(1).strip(); md = md.replace(m.group(0), '', 1)
    b = load_brand(a)
    doc = Document(); page_setup(doc, a.page)
    apply_styles(doc, b, a.numbered)
    lines = md.split('\n'); i = 0; caption = None
    def para(text, style=None):
        p = doc.add_paragraph(style=style); add_inline(p, text); return p
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip(): i += 1; continue
        h = re.match(r'^(#{1,4})\s+(.*)$', ln)
        if h:
            lvl = len(h.group(1)); text = re.sub(TYPED_NO, '', h.group(2)) if a.numbered else h.group(2)
            if re.match(r'(?i)appendix', text): doc.add_paragraph(text, style='Heading 1')
            else: doc.add_paragraph(text, style=f'Heading {lvl}')
            i += 1; continue
        if ln.strip() == '[[revision-history]]': style_table(revision_table(doc, None), b); i += 1; continue
        if ln.strip() == '[[sign-off]]': style_table(signoff_table(doc), b); i += 1; continue
        if ln.strip() in ('---', '***', '[[page-break]]'): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); i += 1; continue
        if ln.startswith('>'):
            block = []
            while i < len(lines) and lines[i].startswith('>'): block.append(lines[i][1:].strip()); i += 1
            text = ' '.join(block); m = SIGNAL_RE.match(text.replace('**', '', 2) if text.startswith('**') else text)
            if m: notice_block(doc, m.group(1), m.group(2))
            else: notice_block(doc, 'NOTE', text)
            continue
        if ln.startswith('Table:'): caption = ln[6:].strip(); i += 1; continue
        if ln.startswith('Figure:'):
            p = doc.add_paragraph(style='Caption'); p.add_run(ln[7:].strip()); i += 1; continue
        if ln.lstrip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c): rows.append(cells)
                i += 1
            if not rows: continue
            if caption:
                cp = doc.add_paragraph(style='Caption'); cp.add_run(caption); caption = None
            ncol = max(len(r) for r in rows); t = doc.add_table(rows=0, cols=ncol)
            for r in rows:
                cells = t.add_row().cells
                for c, v in zip(cells, r + [''] * (ncol - len(r))): add_inline(c.paragraphs[0], v)
            style_table(t, b); doc.add_paragraph().paragraph_format.space_after = Pt(2); continue
        lm = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', ln)
        if lm:
            nested = len(lm.group(1)) >= 2
            style = ('List Number' if lm.group(2)[0].isdigit() else 'List Bullet') + (' 2' if nested else '')
            try: para(lm.group(3), style)
            except KeyError: para(lm.group(3), 'List Bullet')
            i += 1; continue
        buf = [ln.strip()]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|>|\||\s*([-*+]|\d+[.)])\s|Table:|Figure:|\[\[)', lines[i]):
            buf.append(lines[i].strip()); i += 1
        para(' '.join(buf))
    if a.cover: cover_page(doc, meta, b)
    header_footer(doc, meta, b, a.cover)
    finish(doc, meta, b, out)

def finish(doc, meta, b, out):
    cp = doc.core_properties
    cp.title = meta.get('title', '') or ''; cp.subject = meta.get('doc_type', '') or ''
    cp.author = meta.get('author', '') or ''; cp.last_modified_by = meta.get('author', '') or ''
    cp.comments = ''; cp.keywords = meta.get('doc_no', '') or ''; cp.category = meta.get('doc_type', '') or ''
    # company lives in docProps/app.xml: the doc-metadata skill's script owns it
    doc.save(out); print('wrote', out)

# ---------- check ----------
def check(path):
    doc = Document(path); issues = []
    UNIT = r'\d(V|A|W|kW|Hz|mm|kg|°C|°F|psi|bar)(?![A-Za-z])'
    for ti, tb in enumerate(doc.tables):
        bad = [c.text.strip()[:30] for row in tb.rows for c in row.cells if re.search(UNIT, c.text)]
        if bad: issues.append(f'table {ti+1}: value glued to unit in {bad[:3]} -> "20 °C", "2 bar"')
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if re.search(r'\b(see|shown|listed|described)\b[^.]{0,25}\b(above|below)\b', t, re.I): issues.append(f'para {i}: "above/below" cross-reference -> cite section/table number')
        if re.search(UNIT, t): issues.append(f'para {i}: value glued to unit -> "24 V", "65 °C"')
        if p.style.name.startswith('List Number') and t and len(t.split()) > 25: issues.append(f'para {i}: procedural step >25 words')
        if re.match(r'(?i)^(warning|caution|danger|notice)\b', t) and not p._p.getparent().tag.endswith('tc'):
            issues.append(f'para {i}: unboxed signal word -> notice panel')
        if re.search(r'(?i)\b(bottom line|key takeaways?|why it matters)\s*:', t): issues.append(f'para {i}: AI-style label block -> plain prose')
    for s in doc.sections:
        ftxt = ' '.join([p.text for p in s.footer.paragraphs] + [c.text for t in s.footer.tables for row in t.rows for c in row.cells])
        if not ftxt.strip(): issues.append('section without footer identity (doc no / rev / date / page)')
    print('\n'.join(issues) if issues else 'no issues found'); return issues

def main():
    try: sys.stdout.reconfigure(encoding='utf-8')
    except Exception: pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('mode', choices=['brand', 'restyle', 'build', 'check']); ap.add_argument('inp', nargs='?'); ap.add_argument('out', nargs='?')
    ap.add_argument('--brand-file', help='brand profile JSON (default: ./.eng-report/brand.json, then ~/.eng-report/brand.json)')
    ap.add_argument('--plain', action='store_true', help='no brand profile: no logo, neutral accent')
    for k in ('company', 'tagline', 'logo', 'accent', 'head', 'link'): ap.add_argument('--' + k, help='brand mode only')
    ap.add_argument('--no-logo', action='store_true', help='brand mode: drop the stored logo')
    ap.add_argument('--scope', choices=['project', 'user'], default='project', help='brand mode: save in ./.eng-report or ~/.eng-report')
    for k in ('title', 'subtitle', 'doc-type', 'doc-no', 'rev', 'date', 'author', 'classification'): ap.add_argument('--' + k)
    ap.add_argument('--cover', dest='cover', action='store_true', default=None); ap.add_argument('--no-cover', dest='cover', action='store_false')
    ap.add_argument('--numbered', dest='numbered', action='store_true', default=True); ap.add_argument('--no-numbered', dest='numbered', action='store_false')
    ap.add_argument('--page', choices=['letter', 'a4', 'keep'], default='letter')
    a = ap.parse_args()
    if a.mode == 'brand': save_brand(a); return
    if not a.inp: ap.error('input file required')
    if a.mode == 'check': check(a.inp); return
    if not a.out: ap.error('OUT path required')
    if os.path.abspath(a.inp) == os.path.abspath(a.out): ap.error('refusing to overwrite the input; write a new file')
    meta = {k: v for k, v in dict(title=a.title, subtitle=a.subtitle, doc_type=a.doc_type, doc_no=a.doc_no, rev=a.rev,
            date=a.date, author=a.author, classification=a.classification).items() if v}
    if a.cover is None: a.cover = (a.mode == 'build')
    (restyle if a.mode == 'restyle' else build)(a.inp, a.out, meta, a)

if __name__ == '__main__':
    main()
